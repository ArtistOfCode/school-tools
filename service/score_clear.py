import logging
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, Border, PatternFill, Alignment, Protection
from openpyxl.worksheet.worksheet import Worksheet

from model.config_model import Config

grade_match = re.compile(r'([一二三四五六])年级')
header_match = re.compile(r'(姓名|语文|数学|英语)')


class ScoreClearService:
    def __init__(self, config: Config):
        self.config = config

    def clear(self):
        self.rename_file()
        for greade, file in self.config.file_data:
            logging.debug(f'加载成绩数据文件: {greade, file}')
            self.clear_file(greade, file)

    def rename_file(self):
        for file in self.config.data_path.glob('*.xlsx'):
            match = grade_match.search(n1 := file.stem)
            if not match:
                logging.error(f'文件名称不合法: {n1}')
                continue
            if n1 != (n2 := match.group()):
                file.rename(file.parent / f'{n2}.xlsx')
                logging.info(f'文件名称不合法，修改为：{n1} -> {n2}')

    def clear_file(self, greade: str, file: Path):
        wb = load_workbook(file, False, False, True)

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]

            self.pre_handle(sheet)

            self.find_header(sheet)

            self.clear_column(sheet)

            self.clear_row(sheet)

            self.clear_cell(sheet)

            sheet.insert_cols(0, 2)
            sheet.cell(1, 1).value = '年级'
            sheet.cell(1, 2).value = '班级'
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                sheet.cell(row_idx, 1).value = greade
                sheet.cell(row_idx, 2).value = str(sheet_name)

            self.clear_all_styles(sheet)

        wb.save(file)

    @staticmethod
    def pre_handle(sheet: Worksheet):
        # 清除图片
        sheet._images = []
        # 清除合并单元格
        for merged_range in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged_range))

    @staticmethod
    def find_header(sheet: Worksheet):
        for row_idx in range(1, sheet.max_row + 1):
            is_header = False
            for col_idx in range(1, sheet.max_column + 1):
                cell = sheet.cell(row_idx, col_idx).value
                if cell and (is_header := bool(header_match.search(cell))):
                    break
            if is_header:
                if row_idx > 1:
                    logging.debug(f'删除表头前面的行：{sheet.title} --> {row_idx - 1}')
                    sheet.delete_rows(1, row_idx - 1)
                break

    @staticmethod
    def clear_column(sheet: Worksheet):
        delete_cols = []
        for col_idx in range(1, sheet.max_column + 1):
            cell = sheet.cell(1, col_idx).value
            if cell and header_match.search(cell):
                sheet.cell(1, col_idx).value = header_match.search(cell).group()
            else:
                delete_cols.append(col_idx)
        for cols in reversed(delete_cols):
            sheet.delete_cols(cols)

    @staticmethod
    def clear_row(sheet: Worksheet):
        delete_rows = []
        for row_idx, row in enumerate(sheet.iter_rows(2, values_only=True), start=2):
            if not any(row):
                delete_rows.append(row_idx)

        for row in reversed(delete_rows):
            sheet.delete_rows(row)

    @staticmethod
    def clear_cell(sheet: Worksheet):
        for row_idx in range(1, sheet.max_row + 1):
            for col_idx in range(1, sheet.max_column + 1):
                cell = sheet.cell(row_idx, col_idx).value
                if isinstance(cell, str):
                    sheet.cell(row_idx, col_idx).value = cell.strip()

    @staticmethod
    def clear_all_styles(sheet):
        if sheet.auto_filter.ref:
            sheet.auto_filter.ref = None
        if sheet.freeze_panes:
            sheet.freeze_panes = None
        sheet.freeze_panes = 'A2'
        sheet.row_dimensions.clear()
        sheet.column_dimensions.clear()

        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font()
                cell.border = Border()
                cell.fill = PatternFill()
                cell.alignment = Alignment()
                cell.protection = Protection()
                cell.number_format = 'General'
                cell.comment = None
                cell.hyperlink = None
