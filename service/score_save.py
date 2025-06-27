import logging
from typing import List

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pptx import Presentation
from pptx.slide import Slide

from model.config_model import Config
from model.score_model import SubjectScore, ClassScore
from utils.excel_utils import CellIndex, set_cell, set_title_cell, set_float_cell
from utils.ppt_utils import add_layout_slide, set_center_cell, pos, add_table, add_textbox, pos_cm
from utils.utils import Subjects, is_school_class, is_low_grade


class ScoreSave:

    def __init__(self, config: Config):
        self.config = config
        self.grade = None
        self.score: List[ClassScore] = []
        self.excel = Workbook()
        self.excel.remove(self.excel['Sheet'])
        self.ppt = Presentation(self.config.ppt_template_path) if self.config.need_ppt else None

    def add_excel_sheet(self, grade, score: ClassScore):
        self.grade = grade
        self.score = score
        return self.excel.create_sheet(self.grade)

    def write_excel_score(self, sheet: Worksheet):
        row = CellIndex()

        for code, desc in [s.value for s in Subjects]:
            logging.debug(f'当前写入科目: {sheet.title} {desc}')
            _low = is_low_grade(sheet.title)
            if _low and code == Subjects.ENG.code:
                continue

            # 定义表头
            care_score = getattr(self.score[0], code).care_stu_2[0]
            pass_name = '三科' if (not _low and code == Subjects.TWO.code) else ''
            care_name = f'率({care_score:g})' if _low else '平均分'
            headers = ['班级', '总人数', '平均分', f'{pass_name}及格人数', f'{pass_name}及格率', '特优人数', '特优率',
                       f'关爱{care_name}', '总评', '与校平差', '名次']

            # 写入科目名称
            set_title_cell(sheet.cell(row.value, 1), desc)
            row.next()

            # 写入成绩表头
            for _idx, header in enumerate(headers):
                set_title_cell(sheet.cell(row.value, _idx + 1), header)
            row.next()

            # 写入成绩
            for _score in self.score:
                logging.debug(f'当前写入班级: {sheet.title} {desc} {_score.name}')
                _sub_score: SubjectScore = getattr(_score, code)
                _idx = CellIndex(0)
                _cell = lambda: sheet.cell(row.value, _idx.next())

                set_cell(_cell(), _score.name)
                set_cell(_cell(), _score.total_stu)
                set_float_cell(_cell(), _sub_score.mean)
                set_cell(_cell(), _sub_score.pass_stu[0])
                set_float_cell(_cell(), _sub_score.pass_stu[1])
                if code != Subjects.ENG.code:
                    set_cell(_cell(), _sub_score.top_stu[0])
                    set_float_cell(_cell(), _sub_score.top_stu[1])
                if is_low_grade:
                    set_float_cell(_cell(), _sub_score.care_stu_2[2])
                else:
                    set_float_cell(_cell(), _sub_score.care_stu_1[1])
                set_float_cell(_cell(), _sub_score.total)
                if not _score.is_school:
                    set_float_cell(_cell(), _sub_score.diff)
                    set_cell(_cell(), _sub_score.rank)
                row.next()
            row.next()

    def write_excel_care(self, sheet: Worksheet):
        if not self.config.need_care:
            return

        col = CellIndex(12)
        for code, desc in [s.value for s in Subjects]:
            logging.debug(f'当前写入关爱生科目: {sheet.title} {desc}')
            _low = is_low_grade(sheet.title)
            if _low and code == Subjects.ENG.code:
                continue

            _row = CellIndex(0)
            col.next(3)

            for _score in self.score:
                logging.debug(f'当前写入关爱生班级: {sheet.title} {desc} {_score.name}')
                _sub_score: SubjectScore = getattr(_score, code)
                if is_school_class(_score.name) or _sub_score.care_stu_array.size == 0:
                    continue

                _col = CellIndex(col.value)

                # 写入表头
                set_title_cell(sheet.cell(_row.next(), _col.value),
                               f'{_score.name}班{desc}（{_sub_score.care_stu_array.size}）')
                set_title_cell(sheet.cell(_row.next(), _col.value), '姓名')

                _cell = lambda: sheet.cell(_row.value, _col.next())
                if code != Subjects.TWO.code:
                    set_title_cell(_cell(), '分数')
                else:
                    set_title_cell(_cell(), Subjects.CHN.desc)
                    set_title_cell(_cell(), Subjects.MATH.desc)
                    if _score.is_low:
                        set_title_cell(_cell(), Subjects.TWO.desc)
                    else:
                        set_title_cell(_cell(), Subjects.ENG.desc)
                        set_title_cell(_cell(), Subjects.TWO.desc)
                _row.next()

                # 写入关爱生列表
                for stu in _sub_score.care_stu_array:
                    _col = CellIndex(col.value)

                    set_cell(sheet.cell(_row.value, _col.value), stu['name'])
                    if code != Subjects.TWO.code:
                        set_cell(_cell(), stu[code])
                    else:
                        set_cell(_cell(), stu[Subjects.CHN.code])
                        set_cell(_cell(), stu[Subjects.MATH.code])
                        if not _score.is_low:
                            set_cell(_cell(), stu[Subjects.ENG.code])
                        set_cell(_cell(), stu[Subjects.TWO.code])
                    _row.next()
                _row.next()

    def write_pptx(self):
        if not self.config.need_ppt:
            return

        # 创建年级标题页
        add_layout_slide(self.ppt, 1, f'{self.grade}成绩分析')
        logging.debug(f'当前写入PPT年级: {self.grade}')

        for code, desc in [s.value for s in Subjects]:
            logging.debug(f'当前写入PPT科目: {self.grade} {desc}')
            _low = is_low_grade(self.grade)
            if _low and code == Subjects.ENG.code:
                continue

            # 添加科目标题页
            add_layout_slide(self.ppt, 2, f'{desc}情况分析')
            # 添加成绩总结页
            slide = add_layout_slide(self.ppt, 3, f'{desc}情况分析')
            # 添加成绩总结页表格
            self.__add_pptx_table(_low, code, slide)
            # 添加关爱生页
            self.__add_pptx_care(_low, code, desc)

    def save(self):
        self.excel.save(self.config.result_path)
        if self.config.need_ppt:
            self.ppt.save(self.config.ppt_result_path)

    def __add_pptx_table(self, _low, code, slide: Slide):
        # 计算成绩表格表头
        # @formatter:off
        if _low:
            care_score = getattr(self.score[0], code).care_stu_2[0]
            if code in (Subjects.CHN.code, Subjects.MATH.code):
                headers = ['班级', '平均分', '及格率', f'关爱率\v{care_score:g}', '总评', '与校\v平差', '与区\v平差', '名次', '教者']
            else:
                headers = ['班级', '平均分', '及格人数', '及格率', f'关爱率\v{care_score:g}', '总评', '与校\v平差', '与区\v平差', '名次', '班主任']
        else:
            if code in (Subjects.CHN.code, Subjects.MATH.code):
                headers = ['班级', '平均分', '及格率', '关爱\v平均分', '特优率', '总评', '与校\v平差', '与区\v平差', '名次', '教者']
            elif code == Subjects.ENG.code:
                headers = ['班级', '平均分', '及格率', f'关爱\v平均分', '总评', '与校\v平差', '与区\v平差', '名次', '教者']
            else:
                headers = ['班级', '平均分', '三科\v及格人数', f'三科\v及格率', '关爱\v平均分', '总评', '与校\v平差', '与区\v平差', '名次', '班主任']
        # @formatter:on

        # 成绩表格排版
        w, h = 1.2, 0.5
        r, c = len(self.score) + 2, len(headers)
        t = (self.ppt.slide_height.inches - r * h) / 2
        l = (self.ppt.slide_width.inches - c * w) / 2
        table = add_table(slide, (r, c), pos(w, h, t, l))

        for idx, header in enumerate(headers):
            set_center_cell(table.cell(0, idx), header)

        for idx, _score in enumerate(self.score):
            _sub: SubjectScore = getattr(_score, code)

            idx += 1
            col = CellIndex(0)

            _set_cell = lambda s: set_center_cell(table.cell(idx, col.next()), s)
            set_center_cell(table.cell(idx, col.value), _score.name)
            if _low:
                if code in (Subjects.CHN.code, Subjects.MATH.code):
                    _set_cell(to_str(_sub.mean))
                    _set_cell(to_str(_sub.pass_stu[1]))
                    _set_cell(to_str(_sub.care_stu_2[2]))
                else:
                    _set_cell(to_str(_sub.mean))
                    _set_cell(to_str(_sub.pass_stu[0]))
                    _set_cell(to_str(_sub.pass_stu[1]))
                    _set_cell(to_str(_sub.care_stu_2[2]))
            else:
                if code in (Subjects.CHN.code, Subjects.MATH.code):
                    _set_cell(to_str(_sub.mean))
                    _set_cell(to_str(_sub.pass_stu[1]))
                    _set_cell(to_str(_sub.care_stu_1[1]))
                    _set_cell(to_str(_sub.top_stu[1]))
                elif code == Subjects.ENG.code:
                    _set_cell(to_str(_sub.mean))
                    _set_cell(to_str(_sub.pass_stu[1]))
                    _set_cell(to_str(_sub.care_stu_1[1]))
                else:
                    _set_cell(to_str(_sub.mean))
                    _set_cell(to_str(_sub.pass_stu[0]))
                    _set_cell(to_str(_sub.pass_stu[1]))
                    _set_cell(to_str(_sub.care_stu_1[1]))
            _set_cell(to_str(_sub.total))
            if not _score.is_school:
                _set_cell(to_str(_sub.diff))
                _set_cell('')
                _set_cell(str(_sub.rank))
                _set_cell(f'教师{idx}')
        set_center_cell(table.cell(len(self.score) + 1, 0), '区平')

    def __add_pptx_care(self, _low, code, desc):
        if not self.config.need_care:
            return

        if code != Subjects.TWO.code:
            slide = add_layout_slide(self.ppt, 3, f'{desc}关爱生')

            if _low:
                care_score = getattr(self.score[0], code).care_stu_2[0]
                _pos = pos(3, 0.5, self.ppt.slide_height.inches - 1, 1)
                add_textbox(slide, _pos, f'{desc}关爱分数线：{round(care_score, 2):g}')

                for i, _score in enumerate(self.score[:-1]):
                    _sub: SubjectScore = getattr(_score, code)
                    _, care, _ = _sub.care_stu_2

                    l = 2 if i == 0 else i * 5 + 2
                    table = add_table(slide, (care + 2, 2), pos_cm(2.2, 0.9, 3.5, l))

                    row = CellIndex(0)
                    header = table.cell(row.value, 0)
                    header.merge(table.cell(row.value, 1))

                    _size = 14
                    set_center_cell(header, f'{_score.name}班（{care}）', size=_size)
                    row.next()
                    set_center_cell(table.cell(row.value, 0), '姓名', bold=True, size=_size)
                    set_center_cell(table.cell(row.value, 1), '分数', bold=True, size=_size)
                    row.next()

                    for stu in _sub.care_stu_array:
                        set_center_cell(table.cell(row.value, 0), stu['name'], size=_size)
                        set_center_cell(table.cell(row.value, 1), f'{stu[code]:g}', size=_size)
                        row.next()
        else:
            for i, _score in enumerate(self.score[:-1]):
                slide = add_layout_slide(self.ppt, 3, f'{desc}关爱生')

                _sub = getattr(_score, code)
                care_score, care, _ = _sub.care_stu_2
                _pos = pos(2, 0.5, 1.2, 0.8)
                add_textbox(slide, _pos, f'{_score.name}班：{care}个')

                _pos = pos(3, 0.5, self.ppt.slide_height.inches - 1, 1)
                add_textbox(slide, _pos, f'{desc}关爱分数线：{round(care_score, 2):g}')

                if _low:
                    headers = ['姓名', Subjects.CHN.desc, Subjects.MATH.desc, Subjects.TWO.desc]
                else:
                    headers = ['姓名', Subjects.CHN.desc, Subjects.MATH.desc, Subjects.ENG.desc,
                               Subjects.TWO.desc]

                w, h = 5, 1.2
                r, c = care + 1, len(headers)
                t = (self.ppt.slide_height.cm - r * h) / 2
                l = (self.ppt.slide_width.cm - c * w) / 2
                table = add_table(slide, (r, c), pos_cm(w, h, t, l))

                _size = 14
                row = CellIndex(0)
                for _i, h in enumerate(headers):
                    set_center_cell(table.cell(row.value, _i), h, size=_size)
                row.next()

                for stu in _sub.care_stu_array:
                    col = CellIndex(0)
                    set_center_cell(table.cell(row.value, col.value), stu['name'], size=_size)
                    set_center_cell(table.cell(row.value, col.next()), f'{stu[Subjects.CHN.code]:g}', size=_size)
                    set_center_cell(table.cell(row.value, col.next()), f'{stu[Subjects.MATH.code]:g}', size=_size)
                    if not _low:
                        set_center_cell(table.cell(row.value, col.next()), f'{stu[Subjects.ENG.code]:g}',
                                        size=_size)
                    set_center_cell(table.cell(row.value, col.next()), f'{stu[Subjects.TWO.code]:g}', size=_size)
                    row.next()


def to_str(number):
    return f'{round(number, 2):.2f}'
