from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side, Color

__bold = Font(bold=True)
__center = Alignment(horizontal='center')
__fill = PatternFill('solid', fgColor=Color('e3e3e3'))
__thin_side = Side(style='thin', color='000000')
__thin_border = Border(__thin_side, __thin_side, __thin_side, __thin_side)


def set_title_cell(cell: Cell, value) -> Cell:
    return set_cell(cell, value, __bold, __center, __fill)


def set_float_cell(cell: Cell, value) -> Cell:
    cell.number_format = '0.00'
    return set_cell(cell, value)


def set_cell(cell: Cell, value, font: Font = None, align: Alignment = None, fill: PatternFill = None,
             border: Border = __thin_border) -> Cell:
    cell.value = value
    if font is not None: cell.font = font
    if align is not None: cell.alignment = align
    if fill is not None: cell.fill = fill
    if border is not None: cell.border = border
    return cell


class CellIndex:

    def __init__(self, init=1):
        self.value = init

    def next(self, step=1):
        self.value += step
        return self.value
