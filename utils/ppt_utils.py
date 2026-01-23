from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Cm, Pt


def add_layout_slide(ppt, layout, title) -> Slide:
    _layout = ppt.slide_layouts[layout]
    _slide: Slide = ppt.slides.add_slide(_layout)
    _slide.shapes.title.text = title
    return _slide


def add_table(slide: Slide, size, position):
    r, c = size
    w, h, t, l = position
    table = slide.shapes.add_table(r, c, l, t, w, h).table
    for row in table.rows:
        row.height = h
    for col in table.columns:
        col.width = w
    return table


def add_textbox(slide: Slide, position, text):
    w, h, t, l = position
    slide.shapes.add_textbox(l, t, w, h).text_frame.text = text


def set_center_cell(cell, value: str, color='000000', bold=None, size=18):
    cell.text = value
    cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    cell.text_frame.paragraphs[0].font.size = Pt(size)
    cell.text_frame.paragraphs[0].font.bold = bold
    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(color)


def pos(width, height, top, left):
    return Inches(width), Inches(height), Inches(top), Inches(left)


def pos_cm(width, height, top, left):
    return Cm(width), Cm(height), Cm(top), Cm(left)
